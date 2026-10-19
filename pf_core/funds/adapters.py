"""AMC portfolio adapters.

`GenericExcelAdapter` reads the monthly portfolio workbooks AMCs publish. Their layouts vary, so
it does not assume fixed cell positions:

* a **header row** is any row with a cell containing "ISIN" and a cell mentioning market value or
  NAV (case-insensitive). A second header row underneath (e.g. "(Rs. in Lakhs)" under a merged
  "Market Value") is merged into the labels;
* each header row starts a **scheme section**, so one sheet per scheme and several stacked schemes
  per sheet both work. The scheme name and "as on" date are read from the rows just above;
* "Sub Total"/"Total" rows are skipped and **"Grand Total" ends the section**. Stopping at the first
  "Total" would drop the debt and cash lines that V8 needs;
* lines without an ISIN but with a value (TREPS, cash, net receivables) are kept as non-equity.

Anything the adapter had to assume is written to `parse_notes`.
"""

import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Protocol

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from pf_core.funds import config
from pf_core.funds.models import FundHolding, FundPortfolio

Cell = object
Grid = list[list[Cell]]

_ILLUSTRATIVE_MARKERS = ("synthetic", "illustrative")
_DATE_FORMATS = (
    "%d-%b-%Y",
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%d %B %Y",
    "%d %b %Y",
    "%B %d, %Y",
    "%b %d, %Y",
)
_DATE_TEXT = re.compile(
    r"(\d{1,2}[-/ ][A-Za-z]{3,9}[-/ ]\d{4}"  # 31-Mar-2026, 31 March 2026
    r"|\d{1,2}[-/]\d{1,2}[-/]\d{4}"  # 31-03-2026, 31/03/2026
    r"|[A-Za-z]{3,9} \d{1,2}, \d{4})"  # March 31, 2026
)


class AmcAdapter(Protocol):
    """Parses one AMC file into one portfolio per scheme."""

    def parse(self, path: Path) -> list[FundPortfolio]: ...


class GenericExcelAdapter:
    """Header-scanning adapter for AMC `.xlsx` portfolio files."""

    def __init__(self, amc: str | None = None) -> None:
        self.amc = amc

    def parse(self, path: Path) -> list[FundPortfolio]:
        workbook = load_workbook(path, data_only=True)
        amc = self.amc or path.parent.name
        illustrative = any(m in path.name.lower() for m in _ILLUSTRATIVE_MARKERS)
        portfolios: list[FundPortfolio] = []
        for sheet in workbook.worksheets:
            portfolios.extend(_parse_sheet(sheet, path, amc, illustrative))
        return portfolios


def _parse_sheet(sheet: Worksheet, path: Path, amc: str, illustrative: bool) -> list[FundPortfolio]:
    grid: Grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    headers = [i for i, row in enumerate(grid) if _is_header(row)]
    portfolios: list[FundPortfolio] = []
    for n, h in enumerate(headers):
        end = headers[n + 1] if n + 1 < len(headers) else len(grid)
        above_start = max(headers[n - 1] + 1 if n else 0, h - config.HEADER_LOOKBACK_ROWS)
        notes: list[str] = []
        holdings = _section(grid, h, end, notes)
        scheme = _scheme_name(grid[above_start:h]) or sheet.title
        as_of = _as_of(grid[above_start:h]) or _date_from_filename(path, notes)
        portfolios.append(
            FundPortfolio(
                amc=amc,
                scheme_name=scheme,
                as_of_date=as_of,
                source_file=path.name,
                sheet=sheet.title,
                holdings=holdings,
                illustrative=illustrative,
                parse_notes=notes,
            )
        )
    return portfolios


def _text(cell: Cell) -> str:
    return str(cell).strip() if cell is not None else ""


def _is_header(row: list[Cell]) -> bool:
    labels = [_text(c).lower() for c in row]
    return any("isin" in t for t in labels) and any(
        "market" in t or "nav" in t or "net asset" in t for t in labels
    )


def _labels(grid: Grid, h: int) -> tuple[list[str], bool]:
    """Column labels, merging a second header row if the next row holds no data."""
    width = max(len(grid[h]), len(grid[h + 1]) if h + 1 < len(grid) else 0)
    top = [_text(grid[h][c]) if c < len(grid[h]) else "" for c in range(width)]
    if h + 1 >= len(grid):
        return [t.lower() for t in top], False
    nxt = [_text(c) for c in grid[h + 1]] + [""] * width
    second_is_header = not any(re.match(r"^IN[A-Z0-9]{10}$", t) for t in nxt) and not any(
        isinstance(c, int | float) for c in grid[h + 1]
    )
    if second_is_header:
        top = [f"{a} {b}".strip() for a, b in zip(top, nxt[:width], strict=True)]
    return [t.lower() for t in top], second_is_header


def _column(labels: list[str], *needles: str, avoid: tuple[str, ...] = ()) -> int | None:
    for i, label in enumerate(labels):
        if any(n in label for n in needles) and not any(a in label for a in avoid):
            return i
    return None


def _section(grid: Grid, h: int, end: int, notes: list[str]) -> list[FundHolding]:
    labels, two_row = _labels(grid, h)
    isin_col = _column(labels, "isin")
    name_col = _column(labels, "name", "instrument", "company", "security", avoid=("isin",))
    qty_col = _column(labels, "quantity", "qty")
    mv_col = _column(labels, "market", "value", avoid=("nav", "%"))
    pct_col = _column(labels, "nav", "net asset")
    scale = _market_value_scale(labels[mv_col] if mv_col is not None else "", notes)
    holdings: list[FundHolding] = []
    for row in grid[h + (2 if two_row else 1) : end]:
        first = next((_text(c) for c in row if _text(c)), "").lower()
        if re.fullmatch(r"grand\s*total", first):
            break
        if re.match(r"^(sub\s*-?\s*)?total\b", first):
            continue
        isin = _text(_at(row, isin_col)).upper() or None
        mv, pct = _number(_at(row, mv_col)), _number(_at(row, pct_col))
        if isin is None and mv is None and pct is None:
            continue  # a section heading such as "Equity & Equity related"
        holdings.append(
            FundHolding(
                isin=isin,
                name=_text(_at(row, name_col)) or first,
                quantity=_number(_at(row, qty_col)),
                market_value_inr=mv * scale if mv is not None else None,
                pct_nav=pct,
            )
        )
    return holdings


def _at(row: list[Cell], col: int | None) -> Cell:
    return row[col] if col is not None and col < len(row) else None


def _number(cell: Cell) -> float | None:
    """Numbers stay numbers; blanks, dashes and text are missing (None), never zero."""
    if isinstance(cell, bool):
        return None
    if isinstance(cell, int | float):
        return float(cell)
    text = _text(cell).replace(",", "").replace("%", "")
    try:
        return float(text) if text else None
    except ValueError:
        return None


def _market_value_scale(label: str, notes: list[str]) -> float:
    for unit, factor in config.MARKET_VALUE_UNITS.items():
        if unit in label:
            return factor
    notes.append(f"Market value unit not stated in header {label!r}; assumed rupees")
    return 1.0


def _scheme_name(rows: Grid) -> str | None:
    for row in reversed(rows):
        for cell in row:
            text = _text(cell)
            if re.search(r"\b(fund|scheme|plan)\b", text, re.IGNORECASE) and not re.search(
                r"as (on|of|at)", text, re.IGNORECASE
            ):
                return text
    return None


def _as_of(rows: Grid) -> date | None:
    for row in rows:
        line = " ".join(_text(c) for c in row)
        if not re.search(r"as (on|of|at)", line, re.IGNORECASE):
            continue
        for cell in row:
            if isinstance(cell, datetime):
                return cell.date()
            if isinstance(cell, date):
                return cell
        match = _DATE_TEXT.search(line)
        if match:
            return _parse_date(match.group(1))
    return None


def _parse_date(text: str) -> date | None:
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text.replace("/", "-") if "%d-%m" in fmt else text, fmt).date()
        except ValueError:
            continue
    return None


def _date_from_filename(path: Path, notes: list[str]) -> date | None:
    """`<YYYY-MM>.xlsx` → last day of that month, as the drop-folder convention promises."""
    match = re.search(r"(\d{4})-(\d{2})", path.stem)
    if not match:
        notes.append("No 'as on' date found in the sheet or the file name")
        return None
    year, month = int(match.group(1)), int(match.group(2))
    first_next = date(year + month // 12, month % 12 + 1, 1)
    notes.append("as_of_date taken from the file name (month end)")
    return first_next - timedelta(days=1)
