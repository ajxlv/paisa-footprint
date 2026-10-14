"""Build `amc_portfolio_synthetic.xlsx`: an ILLUSTRATIVE AMC monthly-portfolio workbook.

Everything here is made up. Scheme names ("Illustrative ... Fund") are fake, quantities and values
are invented, and ISINs containing "SYN" are deliberately fake. Equity lines reuse the ISINs of the
20 BRSR fixture companies only so the engine has emissions to join against. The workbook imitates
the awkward layouts real AMC files use: title rows, a two-row merged header, values in ₹ lakh,
section sub-totals, cash/TREPS lines without ISINs, a second sheet with two stacked schemes, a
duplicated ISIN, a non-INE ISIN, and a scheme whose weights do not add up to 100%.

Run: python tests/fixtures/synthetic/make_amc_fixture.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

OUT = Path(__file__).with_name("amc_portfolio_synthetic.xlsx")

# (name, ISIN, quantity, market value in ₹ lakh, % to NAV)
LARGE_CAP = [
    ("NMDC Limited", "INE584A01023", 100000, 600.0, 20.0),
    ("HDFC Bank Limited", "INE040A01034", 50000, 900.0, 30.0),
    ("Titan Company Limited", "INE280A01028", 10000, 300.0, 10.0),
    ("Uflex Ltd", "INE516A01017", 20000, 300.0, 10.0),
    ("Varroc Engineering Limited", "INE665L01035", 40000, 300.0, 10.0),
    ("Rajesh Exports Limited", "INE343B01030", 30000, 150.0, 5.0),
]
MID_CAP = [
    ("PRINCE PIPES AND FITTINGS LIMITED", "INE689W01016", 10000, 200.0, 20.0),
    ("GULSHAN POLYOLS LIMITED", "INE255D01024", 20000, 150.0, 15.0),
    ("GULSHAN POLYOLS LIMITED", "INE255D01024", 10000, 100.0, 10.0),  # duplicate ISIN line
    ("Cyient DLM Limited", "INE055S01018", 5000, 250.0, 25.0),
    ("THEMIS MEDICARE LIMITED", "INE083B01024", 8000, 100.0, 10.0),
]
FLEXI_CAP = [
    ("Synthetic Unlisted Co A", "INE0SYN01011", 1000, 500.0, 50.0),
    ("Synthetic Unlisted Co B", "INE0SYN01029", 1000, 300.0, 30.0),
    ("NMDC Limited", "INE584A01023", 10000, 100.0, 10.0),
    ("Odd Instrument", "XX0SYN000001", 10, 50.0, 5.0),  # not an INE ISIN: warning
]


def write_scheme(ws: Worksheet, row: int, scheme: str, equity: list, other: list) -> int:
    """Write one scheme section starting at `row`; return the next free row."""
    ws.cell(row, 1, "ILLUSTRATIVE AMC (synthetic, not real market data)")
    ws.cell(row + 1, 1, scheme)
    ws.cell(row + 2, 1, "Monthly Portfolio Statement as on 31-Mar-2026")
    h = row + 4
    for col, label in enumerate(
        ["Name of the Instrument", "ISIN", "Industry", "Quantity", "Market Value", "% to NAV"], 1
    ):
        ws.cell(h, col, label)
    ws.cell(h + 1, 5, "(Rs. in Lakhs)")  # second header row under a merged market-value header
    for col in (1, 2, 3, 4, 6):
        ws.merge_cells(start_row=h, start_column=col, end_row=h + 1, end_column=col)
    r = h + 2
    ws.cell(r, 1, "Equity & Equity related")
    r += 1
    for name, isin, qty, mv, pct in equity:
        for col, value in enumerate([name, isin, "Synthetic", qty, mv, pct], 1):
            ws.cell(r, col, value)
        r += 1
    ws.cell(r, 1, "Sub Total")
    ws.cell(r, 5, sum(x[3] for x in equity))
    r += 1
    for name, mv, pct in other:
        ws.cell(r, 1, name)
        ws.cell(r, 5, mv)
        ws.cell(r, 6, pct)
        r += 1
    ws.cell(r, 1, "Grand Total")
    r += 3
    return r


def build(path: Path = OUT) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Illustrative Large Cap"
    write_scheme(ws, 1, "Illustrative Large Cap Fund", LARGE_CAP, [("TREPS", 450.0, 15.0)])
    ws2 = wb.create_sheet("Two Schemes")
    nxt = write_scheme(
        ws2, 1, "Illustrative Mid Cap Fund", MID_CAP, [("Net Receivables / (Payables)", 10.0, 1.0)]
    )
    write_scheme(ws2, nxt, "Illustrative Flexi Cap Fund", FLEXI_CAP, [("TREPS", 50.0, 5.0)])
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"Wrote {build()}")
